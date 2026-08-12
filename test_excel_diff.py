"""excel_diff 모듈(개발자 diff 뷰 엑셀 생성) 검증 테스트.

실행: python test_excel_diff.py
"""

import os
import tempfile

import pandas as pd
from openpyxl import load_workbook

from excel_diff import export_diff_xlsx, export_presence_xlsx


def _make_xlsx(path, sheets):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=name, header=False, index=False)


def _fill(cell):
    return cell.fill.fgColor.rgb if (cell.fill and cell.fill.patternType) else None


YELLOW = ("FFFFF2CC", "00FFF2CC")
GREEN = ("FFC6EFCE", "00C6EFCE")
RED = ("FFFFC7CE", "00FFC7CE")


def test_change_insert_delete_alignment():
    """삽입/삭제/변경이 섞여도 diff 정렬로 정확히 구분되는지 검증."""
    with tempfile.TemporaryDirectory() as d:
        a = os.path.join(d, "a.xlsx")
        b = os.path.join(d, "b.xlsx")
        out = os.path.join(d, "out.xlsx")
        _make_xlsx(a, {"S": [["이름", "점수"], ["김철수", 90], ["이영희", 85], ["박민수", 70]]})
        # 철수 점수 변경, 한신입 삽입, 영희 유지, 박민수 삭제
        _make_xlsx(b, {"S": [["이름", "점수"], ["김철수", 95], ["한신입", 60], ["이영희", 85]]})
        result = export_diff_xlsx(a, b, out)

        t = result["total"]
        assert t["changed"] >= 1, "철수 점수 변경이 감지되어야 함"
        assert t["added"] >= 1, "한신입 삽입이 추가로 감지되어야 함"
        assert t["removed"] >= 1, "박민수 삭제가 감지되어야 함"

        wb = load_workbook(out)
        assert wb.sheetnames[0] == "요약"
        ws = wb["S"]
        # 헤더(1행) 아래 어딘가에 초록/빨강 채움이 존재해야 함
        greens = reds = yellows = 0
        for row in ws.iter_rows(min_row=2):
            for c in row:
                f = _fill(c)
                if f in GREEN:
                    greens += 1
                elif f in RED:
                    reds += 1
                elif f in YELLOW:
                    yellows += 1
        assert greens > 0, "삽입 행이 초록으로 칠해져야 함"
        assert reds > 0, "삭제 행이 빨강으로 칠해져야 함"
        assert yellows > 0, "변경 셀이 노랑으로 칠해져야 함"
    print("✓ change/insert/delete alignment & fills")


def test_identical_no_fills():
    with tempfile.TemporaryDirectory() as d:
        a = os.path.join(d, "a.xlsx")
        b = os.path.join(d, "b.xlsx")
        out = os.path.join(d, "out.xlsx")
        data = {"S": [["x", "y"], [1, 2], [3, 4]]}
        _make_xlsx(a, data)
        _make_xlsx(b, data)
        result = export_diff_xlsx(a, b, out)
        assert result["total"]["changed"] == 0
        assert result["total"]["added"] == 0
        assert result["total"]["removed"] == 0

        wb = load_workbook(out)
        ws = wb["S"]
        for row in ws.iter_rows(min_row=2):
            for c in row:
                assert _fill(c) not in YELLOW + GREEN + RED, "동일 파일엔 색이 없어야 함"
    print("✓ identical files produce no diff fills")


def test_sheet_only_in_one():
    with tempfile.TemporaryDirectory() as d:
        a = os.path.join(d, "a.xlsx")
        b = os.path.join(d, "b.xlsx")
        out = os.path.join(d, "out.xlsx")
        _make_xlsx(a, {"공통": [["v"]], "OnlyA": [["삭제"]]})
        _make_xlsx(b, {"공통": [["v"]], "OnlyB": [["추가"]]})
        result = export_diff_xlsx(a, b, out)
        byname = {s["name"]: s for s in result["sheets"]}
        assert byname["OnlyA"]["only_in"] == "A"
        assert byname["OnlyB"]["only_in"] == "B"
        wb = load_workbook(out)
        assert "OnlyA" in wb.sheetnames and "OnlyB" in wb.sheetnames
    print("✓ sheet only in one file")


def test_presence_export_green_red():
    """A의 값이 B에도 있으면 초록, 한쪽에만 있으면 빨강 (양쪽 대칭)."""
    with tempfile.TemporaryDirectory() as d:
        a = os.path.join(d, "a.xlsx")
        b = os.path.join(d, "b.xlsx")
        out = os.path.join(d, "out.xlsx")
        # 사과·바나나는 양쪽에 있음(초록) / 딸기는 A에만, 포도는 B에만(빨강)
        _make_xlsx(a, {"과일": [["사과", "바나나"], ["딸기", None]]})
        _make_xlsx(b, {"과일": [["바나나", "포도"], ["사과", None]]})
        result = export_presence_xlsx(a, b, out)

        # 공통(사과·바나나) → 양쪽 다 초록 = 4칸. 한쪽만(딸기, 포도) = 2칸.
        assert result["total"]["green"] == 4, "공통 값은 양쪽 모두 초록"
        assert result["total"]["red"] == 2, "한쪽에만 있는 값은 빨강"

        wb = load_workbook(out)
        ws = wb["과일"]
        greens, reds = set(), set()
        for row in ws.iter_rows(min_row=2):
            for c in row:
                f = _fill(c)
                if f in GREEN:
                    greens.add(c.value)
                elif f in RED:
                    reds.add(c.value)
        assert greens == {"사과", "바나나"}, f"공통 값이 초록이어야 함: {greens}"
        assert reds == {"딸기", "포도"}, f"한쪽에만 있는 값이 빨강이어야 함: {reds}"
    print("✓ presence export: common=green, one-side-only=red")


if __name__ == "__main__":
    test_change_insert_delete_alignment()
    test_identical_no_fills()
    test_sheet_only_in_one()
    test_presence_export_green_red()
    print("\n모든 테스트 통과 ✅")
